"""Behaviour tests.  Run:  python -m unittest discover -s tests -v

The TF-IDF backend is used so the tests are fast, deterministic and work offline.
"""

import tempfile
import unittest
from pathlib import Path

from doc_assistant import SupportAssistant
from doc_assistant.loader import load_knowledge_base


def new_bot() -> SupportAssistant:
    return SupportAssistant(backend="tfidf")


class TestLoader(unittest.TestCase):
    def test_sections_and_citations(self):
        with tempfile.TemporaryDirectory() as d:
            Path(d, "faq.md").write_text("# My FAQ\n\n## 4.2 Opening Hours\nWe open at 9. We close at 5.\n")
            kb = load_knowledge_base(d)
        self.assertEqual(len(kb.sentences), 2)
        self.assertEqual(kb.sentences[0].citation, "faq.md › §4.2 Opening Hours")

    def test_word_file_with_headings(self):
        import docx
        with tempfile.TemporaryDirectory() as d:
            doc = docx.Document()
            doc.add_heading("Acme Leave Policy", 0)
            doc.add_heading("1 Annual Leave", 1)
            doc.add_paragraph("Employees get 24 days of paid annual leave every year.")
            doc.add_heading("2 Sick Leave", 1)
            doc.add_paragraph("Employees get 12 days of sick leave per year.")
            doc.save(Path(d, "leave.docx"))
            bot = SupportAssistant(docs_dir=d, backend="tfidf")
        self.assertEqual(bot.kb.sentences[0].doc_title, "Acme Leave Policy")
        r = bot.ask("How many sick days do I get?")
        self.assertIn("12 days", r.text)
        self.assertEqual(r.citations, ["leave.docx › §2 Sick Leave"])

    def test_text_file_title_and_numbered_headings(self):
        with tempfile.TemporaryDirectory() as d:
            Path(d, "faq.txt").write_text("Shipping FAQ\n\n1. Delivery Time\nOrders arrive in 3 to 5 days.\n\n"
                                          "2. Returns\nYou can return items within 30 days.\n")
            kb = load_knowledge_base(d)
        self.assertEqual(kb.sentences[0].doc_title, "Shipping FAQ")
        self.assertEqual([s.citation for s in kb.sentences],
                         ["faq.txt › §1 Delivery Time", "faq.txt › §2 Returns"])

    def test_file_without_headings_is_split_into_parts(self):
        with tempfile.TemporaryDirectory() as d:
            Path(d, "notes.txt").write_text(" ".join(f"Fact number {i} is here." for i in range(1, 8)))
            kb = load_knowledge_base(d)
        self.assertEqual(list(dict.fromkeys(s.label for s in kb.sentences)), ["Part 1", "Part 2"])

    def test_empty_folder_raises(self):
        with tempfile.TemporaryDirectory() as d, self.assertRaises(FileNotFoundError):
            load_knowledge_base(d)


class TestConversation(unittest.TestCase):
    def test_every_answer_has_a_citation(self):
        bot = new_bot()
        for q in ["How much does the Plus plan cost?", "Can I get a refund?", "How do I share a folder?"]:
            r = bot.ask(q)
            self.assertTrue(r.citations, q)
            self.assertIn("Source:", r.text)

    def test_correct_section_is_cited(self):
        r = new_bot().ask("Can I get a refund?")
        self.assertIn("§2.3 Refund Policy", r.citations[0])

    def test_never_repeats_a_sentence(self):
        bot = new_bot()
        seen = set()
        for q in ["How much does the Plus plan cost?", "What else do I get if I pay yearly?", "Tell me more"]:
            r = bot.ask(q)
            body = r.text.split("\n\nSource:")[0]
            for sentence in bot.memory.turns[-1].sentences:
                self.assertNotIn(sentence.id, seen, f"repeated: {sentence.text}")
                seen.add(sentence.id)
            self.assertTrue(body)

    def test_follow_up_stays_on_topic(self):
        bot = new_bot()
        bot.ask("My files are not syncing")
        r = bot.ask("Tell me more")
        self.assertEqual(r.kind, "follow_up")
        self.assertIn("§3.2", r.citations[0])
        self.assertIn("Continuing with §3.2", r.text)

    def test_topic_switch_is_announced(self):
        bot = new_bot()
        bot.ask("Can I get a refund?")
        r = bot.ask("How do I turn on two-factor authentication?")
        self.assertIn("Switching topics", r.text)
        self.assertIn("§1.4", r.citations[0])

    def test_repeated_question_gets_short_recap(self):
        bot = new_bot()
        first = bot.ask("How much does the Plus plan cost?")
        bot.ask("Can I get a refund?")
        again = bot.ask("How much does the Plus plan cost?")
        self.assertEqual(again.kind, "repeat")
        self.assertIn(f"turn {first.turn}", again.text)
        self.assertLess(len(again.text), len(first.text) + 120)

    def test_out_of_scope_does_not_guess(self):
        r = new_bot().ask("Do you sell laptops?")
        self.assertEqual(r.kind, "out_of_scope")
        self.assertFalse(r.citations)

    def test_recall_lists_previous_questions(self):
        bot = new_bot()
        bot.ask("Can I get a refund?")
        bot.ask("How do I contact support?")
        r = bot.ask("What have I asked so far?")
        self.assertIn("Can I get a refund?", r.text)
        self.assertIn("How do I contact support?", r.text)

    def test_new_question_with_pronoun_is_not_a_follow_up(self):
        # Worst case for threshold tuning: every pronoun question counts as "vague".
        bot = new_bot()
        bot.th = dict(bot.th, weak_query=1.0)
        bot.ask("How much does the Plus plan cost?")
        r = bot.ask("I deleted a file by mistake, can I recover it?")
        self.assertEqual(r.kind, "answer")
        self.assertIn("Switching topics", r.text)
        self.assertIn("§1.5", r.citations[0])

    def test_every_document_reply_is_cited(self):
        questions = ["How much does the Plus plan cost?", "What else do I get if I pay yearly?",
                     "Can I get a refund?", "How long does it take to get the money back?",
                     "My files are not syncing on my laptop", "Tell me more",
                     "How do I turn on two-factor authentication?", "What if I lose my phone?",
                     "How much does the Plus plan cost?", "I deleted a file by mistake, can I recover it?",
                     "How do I contact support?", "Tell me more", "Tell me more", "Tell me more"]
        bot = new_bot()
        for q in questions:
            r = bot.ask(q)
            if r.kind not in ("smalltalk", "recall", "out_of_scope"):
                self.assertTrue(r.citations, f"turn {r.turn} has no citation: {q}")

    def test_reset_clears_memory(self):
        bot = new_bot()
        bot.ask("Can I get a refund?")
        bot.reset()
        self.assertEqual(bot.memory.turns, [])
        self.assertEqual(bot.ask("Can I get a refund?").kind, "answer")


if __name__ == "__main__":
    unittest.main()
